import sqlite3
import os
import openpyxl

DATABASE_FILENAME = "university.db"

class DatabaseManagement:
    def __init__(self, db_path: str = None):
        base = os.path.dirname(os.path.abspath(__file__))
        if db_path:
            self.db_path = db_path
        else:
            data_dir = os.path.join(base, "data")
            os.makedirs(data_dir, exist_ok=True)
            self.db_path = os.path.join(data_dir, DATABASE_FILENAME)

        self.conn = sqlite3.connect(self.db_path)
        self.conn.row_factory = sqlite3.Row
        self._create_tables()
        self._seed_data()

    def _create_tables(self):
        c = self.conn.cursor()
        c.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id TEXT UNIQUE,
                name TEXT,
                email TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                role TEXT NOT NULL DEFAULT 'student',
                semester INTEGER DEFAULT 1
            );

            CREATE TABLE IF NOT EXISTS courses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                code TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL,
                credit INTEGER DEFAULT 3,
                semesters INTEGER NOT NULL,
                p_code TEXT,
                day TEXT,
                t_start TEXT,
                t_end TEXT,
                capacity INTEGER DEFAULT 35
            );

            CREATE TABLE IF NOT EXISTS enrollments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL REFERENCES users(id),
                course_id INTEGER NOT NULL REFERENCES courses(id),
                semester INTEGER NOT NULL,
                status TEXT NOT NULL DEFAULT 'pending',
                grade TEXT DEFAULT NULL,
                grade_point REAL DEFAULT NULL,
                UNIQUE(student_id, course_id)
            );
        """)
        self.conn.commit()

    def _schedule(self, t):
        if not t:
            return ("", "", "")
        parts = str(t).split()
        if len(parts) >= 2:
            day = parts[0]
            times = parts[1].split("-")
            ts = times[0] if len(times) > 0 else ""
            te = times[1] if len(times) > 1 else ""
            return (day, ts, te)
        return (str(t), "", "")

    def _seed_data(self):
        # 1. Ensure Default Administrator Exists
        self._exec(
            "INSERT OR IGNORE INTO users (email, password, role, student_id, name) VALUES (?,?,?,?,?)",
            ("admin@university.edu", "admin", "admin", "admin", "System Administrator")
        )
        self.conn.commit()

        # If already populated with students, return
        if self._scalar("SELECT COUNT(*) FROM users") > 1:
            return

        base = os.path.dirname(os.path.abspath(__file__))

        # 2. Seed Users from Excel
        user_excel_paths = [
            os.path.join(base, "data", "Users.xlsx"),
            os.path.join(base, "Users.xlsx")
        ]
        user_path = next((p for p in user_excel_paths if os.path.exists(p)), None)

        if user_path:
            try:
                wb = openpyxl.load_workbook(user_path)
                ws = wb.active
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue
                    sid = str(row[0]).strip()
                    name = str(row[1]).strip() if len(row) > 1 and row[1] else "Student"
                    email = str(row[2]).strip() if len(row) > 2 and row[2] else f"{sid}@university.edu"

                    if sid == "admin":
                        continue
                    password = sid[-5:] if len(sid) >= 5 else "12345"
                    # Default semester: Alex Morgan starts in semester 2 with transcript
                    sem = 2 if sid in ("202401001", "202401003") else 1
                    self._exec(
                        "INSERT OR IGNORE INTO users (student_id, name, email, password, role, semester) VALUES (?,?,?,?,?,?)",
                        (sid, name, email, password, "student", sem)
                    )
            except Exception as e:
                print(f"User seed error: {e}")

        # 3. Seed Courses from Excel
        course_excel_paths = [
            os.path.join(base, "data", "Courses.xlsx"),
            os.path.join(base, "Courses.xlsx"),
            os.path.join(base, "data", "Lessons.xlsx"),
            os.path.join(base, "Lessons.xlsx")
        ]
        course_path = next((p for p in course_excel_paths if os.path.exists(p)), None)

        if course_path:
            try:
                wb = openpyxl.load_workbook(course_path)
                ws = wb.active
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not row[1]:
                        continue
                    sem_str = str(row[0]).strip() if row[0] else "1"
                    code = str(row[1]).strip()
                    if code in ("Code", "Semester"):
                        continue
                    name = str(row[2]).strip()
                    try:
                        credit = int(row[3]) if row[3] else 4
                    except Exception:
                        credit = 4
                    p = str(row[4]).strip() if row[4] else "-"
                    p_code = None if p in ("-", "", "None") else p
                    t = str(row[5]).strip() if row[5] else ""
                    cap = int(row[6]) if len(row) > 6 and row[6] else 35

                    sem = 2 if "II" in sem_str or "2" in sem_str else 1
                    day, ts, te = self._schedule(t)

                    self._exec(
                        "INSERT OR IGNORE INTO courses (code, name, credit, semesters, p_code, day, t_start, t_end, capacity) "
                        "VALUES (?,?,?,?,?,?,?,?,?)",
                        (code, name, credit, sem, p_code, day, ts, te, cap)
                    )
            except Exception as e:
                print(f"Course seed error: {e}")

        self.conn.commit()

        # 4. Seed Realistic Grades and Completed Transcripts
        # For Alex Morgan (202401001), seed completed Semester 1 with strong grades
        alex = self.fetchone("SELECT id FROM users WHERE student_id='202401001'")
        if alex:
            sem1_courses = self.fetchall("SELECT id, code FROM courses WHERE semesters=1")
            grades = {
                "CS101": ("AA", 4.0),
                "MATH101": ("BA", 3.5),
                "PHYS101": ("AA", 4.0),
                "ENG101": ("BA", 3.5),
                "MATH103": ("BB", 3.0),
                "HIST101": ("AA", 4.0)
            }
            for c in sem1_courses:
                g_letter, g_pt = grades.get(c["code"], ("BB", 3.0))
                self._exec(
                    "INSERT OR IGNORE INTO enrollments (student_id, course_id, semester, status, grade, grade_point) "
                    "VALUES (?, ?, 1, 'approved', ?, ?)",
                    (alex["id"], c["id"], g_letter, g_pt)
                )

        # For Taylor Bennett (202401003), seed mixed grades (including one FF to demonstrate probation & retake)
        taylor = self.fetchone("SELECT id FROM users WHERE student_id='202401003'")
        if taylor:
            sem1_courses = self.fetchall("SELECT id, code FROM courses WHERE semesters=1")
            t_grades = {
                "CS101": ("CC", 2.0),
                "MATH101": ("FF", 0.0),  # Failed
                "PHYS101": ("DC", 1.5),  # Conditional
                "ENG101": ("CB", 2.5),
                "MATH103": ("DD", 1.0),
                "HIST101": ("CC", 2.0)
            }
            for c in sem1_courses:
                g_letter, g_pt = t_grades.get(c["code"], ("CC", 2.0))
                self._exec(
                    "INSERT OR IGNORE INTO enrollments (student_id, course_id, semester, status, grade, grade_point) "
                    "VALUES (?, ?, 1, 'approved', ?, ?)",
                    (taylor["id"], c["id"], g_letter, g_pt)
                )

        # For Jordan Lee (202401002), seed fresh pending enrollment requests
        jordan = self.fetchone("SELECT id FROM users WHERE student_id='202401002'")
        if jordan:
            sem1_courses = self.fetchall("SELECT id, code FROM courses WHERE semesters=1 LIMIT 3")
            for c in sem1_courses:
                self._exec(
                    "INSERT OR IGNORE INTO enrollments (student_id, course_id, semester, status) "
                    "VALUES (?, ?, 1, 'pending')",
                    (jordan["id"], c["id"])
                )

        self.conn.commit()

    def _exec(self, sql, params=()):
        return self.conn.execute(sql, params)

    def _scalar(self, sql, params=()):
        res = self.conn.execute(sql, params).fetchone()
        return res[0] if res else 0

    def fetchone(self, sql, params=()):
        return self.conn.execute(sql, params).fetchone()

    def fetchall(self, sql, params=()):
        return self.conn.execute(sql, params).fetchall()

    def execute(self, sql, params=()):
        cur = self.conn.execute(sql, params)
        self.conn.commit()
        return cur
